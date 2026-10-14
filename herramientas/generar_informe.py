# Arma el informe en PDF y Word.
import re
from decimal import Decimal
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import mysql.connector
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, Image, NextPageTemplate, PageBreak,
                                PageTemplate, Paragraph, Preformatted, Spacer, Table,
                                TableStyle, KeepTogether)
from reportlab.platypus.tableofcontents import TableOfContents

RAIZ = Path(__file__).resolve().parent.parent
HERR = RAIZ / "herramientas"
SALIDA = HERR / "salida"
SCRIPT_SQL = RAIZ / "TP3_Vinos_Lucas_Llano.sql"
CONTENIDO = HERR / "informe.md"
PDF = RAIZ / "Informe_TP3_Lucas_Llano.pdf"
DOCX = RAIZ / "Informe_TP3_Lucas_Llano.docx"

TITULO = "Base de datos relacional de reseñas de vinos"
SUBTITULO = "Diseño, implementación en MySQL y análisis multidimensional"
AUTOR = "Lucas Llano"
FECHA = "14 de octubre de 2026"
ENCABEZADO = "TP3 · Base de datos relacional de reseñas de vinos"

AZUL = "#2a78d6"
TINTA, TINTA2, MUDO, GRILLA = "#0b0b0b", "#52514e", "#898781", "#e4e3df"


# busca los comentarios '-- C01 - ...' del script y se queda con la consulta de abajo
def leer_consultas():
    lineas = SCRIPT_SQL.read_text(encoding="utf-8").splitlines()
    consultas, i = {}, 0
    while i < len(lineas):
        m = re.match(r"-- ([CM]\d{2}) - ", lineas[i])
        i += 1
        if not m:
            continue
        while lineas[i].startswith("--"):
            i += 1
        cuerpo = []
        while True:
            cuerpo.append(lineas[i])
            i += 1
            if cuerpo[-1].rstrip().endswith(";"):
                break
        consultas[m.group(1)] = "\n".join(cuerpo)
    return consultas


def ejecutar(cur, sql):
    cur.execute(sql.strip().rstrip(";"))
    filas = cur.fetchall()
    columnas = [d[0] for d in cur.description]
    while cur.nextset():  # el CALL devuelve un resultado de más
        pass
    return columnas, filas


def formato(v):
    if v is None:
        return "NULL"
    if isinstance(v, (Decimal, float)):
        return f"{v}".replace(".", ",")
    if isinstance(v, bytes):
        return v.decode("utf-8")
    return str(v)


# separa informe.md en bloques: títulos, párrafos, tablas, código y las
# marcas ::consulta, ::figura y ::nota
def parsear(texto):
    bloques, lineas, i = [], texto.splitlines(), 0
    while i < len(lineas):
        l = lineas[i]
        if not l.strip():
            i += 1
        elif l.startswith("```"):
            tipo, cuerpo = l[3:].strip(), []
            i += 1
            while not lineas[i].startswith("```"):
                cuerpo.append(lineas[i])
                i += 1
            i += 1
            bloques.append((tipo, "\n".join(cuerpo)))
        elif l.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            bloques.append(("tabla", filas))
        elif l.startswith("## "):
            bloques.append(("h2", l[3:].strip())); i += 1
        elif l.startswith("# "):
            bloques.append(("h1", l[2:].strip())); i += 1
        elif l.startswith("::"):
            nombre, _, resto = l[2:].partition(" ")
            bloques.append((nombre, resto.strip())); i += 1
        else:
            parrafo = []
            while i < len(lineas) and lineas[i].strip() and not lineas[i].startswith(("#", "::", "|", "```")):
                parrafo.append(lineas[i].strip())
                i += 1
            bloques.append(("p", " ".join(parrafo)))
    return bloques


# reemplaza cada ::consulta por el código + la tabla con el resultado
def resolver(bloques, consultas, cur):
    resultados, salida = {}, []
    for tipo, dato in bloques:
        if tipo == "consulta":
            ident, maximo = dato.split()
            columnas, filas = ejecutar(cur, consultas[ident])
            resultados[ident] = (columnas, filas)
            salida.append(("codigo", consultas[ident]))
            visibles = filas[:int(maximo)]
            tabla = [columnas] + [[formato(v) for v in f] for f in visibles]
            nota = f"Resultado de {ident}"
            nota += f" (primeras {len(visibles)} de {len(filas)} filas)." if len(filas) > len(visibles) else f" ({len(filas)} filas)."
            salida.append(("tabla", tabla))
            salida.append(("nota", nota))
        elif tipo == "ejecutar":
            columnas, filas = ejecutar(cur, dato)
            salida.append(("tabla", [columnas] + [[formato(v) for v in f] for f in filas]))
        elif tipo == "sql":
            salida.append(("codigo", dato))
        else:
            salida.append((tipo, dato))
    return salida, resultados


def estilo_ejes(ax):
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    for lado in ("left", "bottom"):
        ax.spines[lado].set_color(MUDO)
    ax.tick_params(colors=TINTA2, labelsize=9)
    ax.grid(color=GRILLA, linewidth=0.6)
    ax.set_axisbelow(True)


def graficos(res):
    SALIDA.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "Arial", "text.color": TINTA,
                         "axes.labelcolor": TINTA2, "figure.dpi": 200})

    # Reseñas por país (C01)
    cols, filas = res["C01"]
    paises = [f[0] for f in filas][::-1]
    valores = [f[cols.index("resenas")] for f in filas][::-1]
    fig, ax = plt.subplots(figsize=(7, 4.6))
    ax.barh(paises, valores, color=AZUL, height=0.65)
    estilo_ejes(ax)
    ax.grid(axis="y", visible=False)
    for y, v in enumerate(valores):
        ax.text(v + 400, y, f"{v:,}".replace(",", "."), va="center", fontsize=8, color=TINTA2)
    ax.set_xlabel("Reseñas")
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: f"{x:,.0f}".replace(",", ".")))
    fig.tight_layout()
    fig.savefig(SALIDA / "fig_paises.png")
    plt.close(fig)

    # Precio promedio por puntos (C04)
    cols, filas = res["C04"]
    puntos = [f[0] for f in filas]
    precio = [float(f[cols.index("precio_promedio")]) for f in filas]
    fig, ax = plt.subplots(figsize=(7, 3.8))
    ax.plot(puntos, precio, color=AZUL, linewidth=2, marker="o", markersize=4)
    ax.set_yscale("log")
    ax.set_yticks([15, 25, 50, 100, 200, 500])
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:.0f}"))
    ax.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    ax.set_xticks(puntos[::2])
    estilo_ejes(ax)
    ax.set_xlabel("Puntos")
    ax.set_ylabel("Precio promedio (USD, escala log.)")
    for p, v in zip(puntos, precio):
        if p in (80, 90, 100):
            ax.annotate(formato(Decimal(str(v))), (p, v), textcoords="offset points",
                        xytext=(-6, 8), fontsize=8, color=TINTA2, ha="right")
    fig.tight_layout()
    fig.savefig(SALIDA / "fig_precio_puntos.png")
    plt.close(fig)

    # Categorías de puntaje por país (M03), barras apiladas
    cols, filas = res["M03"]
    filas = list(filas)[::-1]
    nombres = ["Aceptable", "Bueno", "Muy bueno", "Excelente", "Soberbio o más"]
    rampa = ["#cde2fb", "#86b6ef", "#3987e5", "#1c5cab", "#0d366b"]
    fig, ax = plt.subplots(figsize=(7, 4.4))
    izquierda = [0.0] * len(filas)
    for j, (nombre, color) in enumerate(zip(nombres, rampa)):
        anchos = [float(f[2 + j]) for f in filas]
        ax.barh([f[0] for f in filas], anchos, left=izquierda, color=color,
                edgecolor="white", linewidth=1, height=0.7, label=nombre)
        izquierda = [a + b for a, b in zip(izquierda, anchos)]
    estilo_ejes(ax)
    ax.grid(visible=False)
    ax.set_xlim(0, 100)
    ax.set_xlabel("% de reseñas del país")
    ax.legend(ncol=5, fontsize=7.5, frameon=False, loc="lower center", bbox_to_anchor=(0.45, 1.0))
    fig.tight_layout()
    fig.savefig(SALIDA / "fig_categorias.png")
    plt.close(fig)

    # Mapa de calor país x rango de precio (M07)
    cols, filas = res["M07"]
    matriz = [[float(v) for v in f[1:]] for f in filas]
    fig, ax = plt.subplots(figsize=(7, 4.6))
    im = ax.imshow(matriz, cmap="Blues", aspect="auto", vmin=83, vmax=95)
    ax.set_xticks(range(5), ["< 15", "15-29", "30-59", "60-99", ">= 100"])
    ax.set_yticks(range(len(filas)), [f[0] for f in filas])
    ax.set_xlabel("Rango de precio (USD)")
    ax.tick_params(colors=TINTA2, labelsize=9, length=0)
    for s in ax.spines.values():
        s.set_visible(False)
    for i, fila in enumerate(matriz):
        for j, v in enumerate(fila):
            ax.text(j, i, f"{v:.2f}".replace(".", ","), ha="center", va="center", fontsize=8,
                    color="white" if v > 90 else TINTA)
    barra = fig.colorbar(im, ax=ax, shrink=0.8)
    barra.set_label("Puntaje promedio", color=TINTA2)
    barra.outline.set_visible(False)
    fig.tight_layout()
    fig.savefig(SALIDA / "fig_calor.png")
    plt.close(fig)


FUENTES = Path("C:/Windows/Fonts")
pdfmetrics.registerFont(TTFont("TNR", str(FUENTES / "times.ttf")))
pdfmetrics.registerFont(TTFont("TNR-B", str(FUENTES / "timesbd.ttf")))
pdfmetrics.registerFont(TTFont("TNR-I", str(FUENTES / "timesi.ttf")))
pdfmetrics.registerFont(TTFont("TNR-BI", str(FUENTES / "timesbi.ttf")))
pdfmetrics.registerFont(TTFont("Consolas", str(FUENTES / "consola.ttf")))
pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-B", italic="TNR-I", boldItalic="TNR-BI")

E = {
    "cuerpo": ParagraphStyle("cuerpo", fontName="TNR", fontSize=12, leading=18,
                             firstLineIndent=1.27 * cm, alignment=TA_JUSTIFY, spaceAfter=6),
    "h1": ParagraphStyle("h1", fontName="TNR-B", fontSize=14, leading=21,
                         spaceBefore=6, spaceAfter=10, keepWithNext=1),
    "h2": ParagraphStyle("h2", fontName="TNR-I", fontSize=12, leading=18,
                         spaceBefore=10, spaceAfter=4, keepWithNext=1),
    "codigo": ParagraphStyle("codigo", fontName="Consolas", fontSize=7.5, leading=9.4,
                             spaceBefore=8, spaceAfter=12),
    "celda": ParagraphStyle("celda", fontName="TNR", fontSize=9, leading=10.8),
    "celda_h": ParagraphStyle("celda_h", fontName="TNR-B", fontSize=9, leading=10.8),
    "nota": ParagraphStyle("nota", fontName="TNR-I", fontSize=10, leading=13,
                           alignment=TA_CENTER, spaceBefore=3, spaceAfter=12),
    "indice_t": ParagraphStyle("indice_t", fontName="TNR-B", fontSize=14, leading=21, spaceAfter=14),
}
TOC_ESTILOS = [
    ParagraphStyle("toc1", fontName="TNR-B", fontSize=12, leading=18, leftIndent=0, firstLineIndent=0),
    ParagraphStyle("toc2", fontName="TNR-I", fontSize=12, leading=18, leftIndent=0.8 * cm, firstLineIndent=0),
]


# Preformatted no pinta el fondo, así que se dibuja el rectángulo a mano
class Codigo(Preformatted):
    def draw(self):
        self.canv.saveState()
        self.canv.setFillColor(colors.HexColor("#f3f2ef"))
        self.canv.rect(-5, -4, self.width + 10, self.height + 6, stroke=0, fill=1)
        self.canv.restoreState()
        super().draw()


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class Documento(BaseDocTemplate):
    def __init__(self, ruta):
        super().__init__(str(ruta), pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                         topMargin=2.5 * cm, bottomMargin=2.5 * cm, title=TITULO, author=AUTOR,
                         subject="Trabajo Práctico 3")
        vertical = Frame(2.5 * cm, 2.5 * cm, A4[0] - 5 * cm, A4[1] - 5 * cm, id="v")
        self.addPageTemplates([
            PageTemplate("portada", [vertical]),
            PageTemplate("normal", [vertical], onPage=self.marco),
        ])

    @staticmethod
    def marco(canv, doc):
        ancho, alto = canv._pagesize
        canv.saveState()
        canv.setFont("TNR-I", 9)
        canv.setFillColor(colors.HexColor(TINTA2))
        canv.drawString(2.5 * cm, alto - 1.6 * cm, ENCABEZADO)
        canv.setStrokeColor(colors.HexColor("#bdbcb7"))
        canv.setLineWidth(0.5)
        canv.line(2.5 * cm, alto - 1.75 * cm, ancho - 2.5 * cm, alto - 1.75 * cm)
        canv.line(2.5 * cm, 1.75 * cm, ancho - 2.5 * cm, 1.75 * cm)
        canv.setFont("TNR", 9)
        canv.drawString(2.5 * cm, 1.3 * cm, f"{AUTOR} · {FECHA}")
        canv.drawRightString(ancho - 2.5 * cm, 1.3 * cm, f"Página {doc.page - 1}")
        canv.restoreState()

    def afterFlowable(self, f):
        nivel = getattr(f, "nivel_toc", None)
        if nivel is None:
            return
        texto = f.getPlainText()
        clave = f"s{self.seq.nextf('sec')}"
        self.canv.bookmarkPage(clave)
        self.canv.addOutlineEntry(texto, clave, level=nivel)
        self.notify("TOCEntry", (nivel, texto, self.page - 1, clave))


def tabla_pdf(filas, ancho_total):
    # cada columna tiene como mínimo el ancho de su palabra más larga (si no
    # corta palabras al medio) y lo que sobra se reparte
    n = len(filas[0])
    medir = lambda t, i: pdfmetrics.stringWidth(t, "TNR-B" if i == 0 else "TNR", 9) + 9
    minimos = [max(medir(w, i) for i, f in enumerate(filas) for w in str(f[j]).split() or [""])
               for j in range(n)]
    deseados = [max(medir(str(f[j]), i) for i, f in enumerate(filas)) for j in range(n)]
    if sum(deseados) <= ancho_total:
        anchos = deseados
    else:
        sobra = max(ancho_total - sum(minimos), 0)
        faltan = [d - m for d, m in zip(deseados, minimos)]
        anchos = [m + sobra * f / (sum(faltan) or 1) for m, f in zip(minimos, faltan)]
        escala = ancho_total / sum(anchos)
        anchos = [a * min(escala, 1) for a in anchos]
    datos = [[Paragraph(esc(str(c)), E["celda_h" if i == 0 else "celda"]) for c in f] for i, f in enumerate(filas)]
    t = Table(datos, colWidths=anchos, repeatRows=1, hAlign="CENTER")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eefa")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b5b4ae")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def imagen(ruta, ancho_cm):
    from reportlab.lib.utils import ImageReader
    w, h = ImageReader(str(ruta)).getSize()
    ancho = ancho_cm * cm
    return Image(str(ruta), width=ancho, height=ancho * h / w)


def portada_pdf():
    return [
        Spacer(1, 3.2 * cm),
        Paragraph("Trabajo Práctico N.º 3", ParagraphStyle("p0", fontName="TNR", fontSize=14, alignment=TA_CENTER, leading=20)),
        Spacer(1, 1.6 * cm),
        Paragraph(TITULO, ParagraphStyle("p1", fontName="TNR-B", fontSize=22, alignment=TA_CENTER, leading=28)),
        Spacer(1, 0.5 * cm),
        Paragraph(SUBTITULO, ParagraphStyle("p2", fontName="TNR-I", fontSize=14, alignment=TA_CENTER, leading=20)),
        Spacer(1, 1.2 * cm),
        Paragraph("Conjunto de datos: winemag-data-130k-v2 (WineEnthusiast)",
                  ParagraphStyle("p3", fontName="TNR", fontSize=12, alignment=TA_CENTER, leading=18)),
        Spacer(1, 7 * cm),
        Paragraph(f"<b>Autor:</b> {AUTOR}", ParagraphStyle("p4", fontName="TNR", fontSize=13, alignment=TA_CENTER, leading=20)),
        Paragraph(f"<b>Fecha:</b> {FECHA}", ParagraphStyle("p5", fontName="TNR", fontSize=13, alignment=TA_CENTER, leading=20)),
    ]


def generar_pdf(bloques):
    doc = Documento(PDF)
    ancho = doc.width
    historia = portada_pdf() + [NextPageTemplate("normal"), PageBreak()]
    toc = TableOfContents(levelStyles=TOC_ESTILOS, dotsMinLevel=0)
    historia += [Paragraph("Índice", E["indice_t"]), toc]
    pendiente = None  # tabla que espera su nota, para que no queden separadas
    for tipo, dato in bloques:
        if tipo == "h1":
            historia.append(PageBreak())
            p = Paragraph(esc(dato), E["h1"]); p.nivel_toc = 0
            historia.append(p)
        elif tipo == "h2":
            p = Paragraph(esc(dato), E["h2"]); p.nivel_toc = 1
            historia.append(p)
        elif tipo == "p":
            historia.append(Paragraph(esc(dato), E["cuerpo"]))
        elif tipo == "codigo":
            historia.append(Codigo(dato, E["codigo"], maxLineLength=112, newLineChars="    "))
        elif tipo == "tabla":
            pendiente = tabla_pdf(dato, ancho)
            if len(dato) > 26:  # las largas sí se pueden cortar entre páginas
                historia.append(pendiente)
                pendiente = None
        elif tipo == "nota":
            nota = Paragraph(esc(dato), E["nota"])
            historia.append(KeepTogether([pendiente, nota]) if pendiente else nota)
            pendiente = None
        elif tipo == "figura":
            archivo, leyenda, ancho_cm = [x.strip() for x in dato.split("|")]
            historia.append(KeepTogether([Spacer(1, 4), imagen(SALIDA / archivo, float(ancho_cm)),
                                          Paragraph(esc(leyenda), E["nota"])]))
    doc.multiBuild(historia)


# campos de Word (número de página, índice). python-docx no los trae
def campo(parrafo, instruccion, relleno=""):
    def elem(tipo, **attrs):
        e = OxmlElement(tipo)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e
    r1 = parrafo.add_run()._r
    r1.append(elem("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = parrafo.add_run()._r
    instr = elem("w:instrText", **{"xml:space": "preserve"})
    instr.text = instruccion
    r2.append(instr)
    r3 = parrafo.add_run()._r
    r3.append(elem("w:fldChar", **{"w:fldCharType": "separate"}))
    parrafo.add_run(relleno)
    r4 = parrafo.add_run()._r
    r4.append(elem("w:fldChar", **{"w:fldCharType": "end"}))


def fuente(estilo, nombre, tam, negrita=False, cursiva=False):
    estilo.font.name = nombre
    estilo.font.size = Pt(tam)
    estilo.font.bold = negrita
    estilo.font.italic = cursiva
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), nombre)
    estilo.element.rPr.rFonts.set(qn("w:hAnsi"), nombre)


def sombrear(celda_o_parrafo, color, es_parrafo=False):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    destino = celda_o_parrafo._p.get_or_add_pPr() if es_parrafo else celda_o_parrafo._tc.get_or_add_tcPr()
    destino.append(sh)


def margenes(seccion):
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(seccion, lado, Cm(2.5))
    seccion.header_distance = Cm(1.2)
    seccion.footer_distance = Cm(1.2)


def generar_docx(bloques):
    d = Document()
    st = d.styles
    fuente(st["Normal"], "Times New Roman", 12)
    pf = st["Normal"].paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.27)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    for nombre, tam, neg, cur in (("Heading 1", 14, True, False), ("Heading 2", 12, False, True)):
        fuente(st[nombre], "Times New Roman", tam, neg, cur)
        st[nombre].paragraph_format.first_line_indent = Cm(0)
        st[nombre].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st[nombre].paragraph_format.space_before = Pt(12 if tam == 14 else 10)
        st[nombre].paragraph_format.space_after = Pt(8 if tam == 14 else 4)
    codigo = st.add_style("Codigo", 1)
    fuente(codigo, "Consolas", 7.5)
    codigo.paragraph_format.line_spacing = 1.0
    codigo.paragraph_format.first_line_indent = Cm(0)
    codigo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    codigo.paragraph_format.space_after = Pt(0)
    nota = st.add_style("Nota", 1)
    fuente(nota, "Times New Roman", 10, cursiva=True)
    nota.paragraph_format.first_line_indent = Cm(0)
    nota.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nota.paragraph_format.line_spacing = 1.0
    nota.paragraph_format.space_after = Pt(12)

    # portada (sin encabezado ni pie)
    s1 = d.sections[0]
    margenes(s1)

    def centrado(texto, tam, neg=False, cur=False, antes=0):
        p = d.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        r = p.add_run(texto)
        r.font.size, r.bold, r.italic = Pt(tam), neg, cur
        return p

    centrado("Trabajo Práctico N.º 3", 14, antes=90)
    centrado(TITULO, 22, neg=True, antes=40)
    centrado(SUBTITULO, 14, cur=True, antes=10)
    centrado("Conjunto de datos: winemag-data-130k-v2 (WineEnthusiast)", 12, antes=30)
    centrado(f"Autor: {AUTOR}", 13, antes=190)
    centrado(f"Fecha: {FECHA}", 13)

    # resto del documento, la numeración empieza en 1 en el índice
    s2 = d.add_section(WD_SECTION.NEW_PAGE)
    margenes(s2)
    for parte in (s2.header, s2.footer):
        parte.is_linked_to_previous = False
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    s2._sectPr.append(pg)

    def marco(seccion):
        h = seccion.header.paragraphs[0]
        h.paragraph_format.first_line_indent = Cm(0)
        r = h.add_run(ENCABEZADO)
        r.italic, r.font.size = True, Pt(9)
        f = seccion.footer.paragraphs[0]
        f.paragraph_format.first_line_indent = Cm(0)
        f.add_run(f"{AUTOR} · {FECHA}    |    Página ").font.size = Pt(9)
        campo(f, "PAGE", "1")
    marco(s2)

    p = d.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Índice")
    r.bold, r.font.size = True, Pt(14)
    t = d.add_paragraph()
    t.paragraph_format.first_line_indent = Cm(0)
    campo(t, 'TOC \\o "1-2" \\h \\z \\u')
    ajustes = d.settings.element
    actualizar = OxmlElement("w:updateFields")
    actualizar.set(qn("w:val"), "true")
    ajustes.append(actualizar)

    for tipo, dato in bloques:
        if tipo == "h1":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            d.add_heading(dato, level=1)
        elif tipo == "h2":
            d.add_heading(dato, level=2)
        elif tipo == "p":
            d.add_paragraph(dato)
        elif tipo == "codigo":
            p = d.add_paragraph(style="Codigo")
            lineas = dato.split("\n")
            for k, linea in enumerate(lineas):
                p.add_run(linea)
                if k < len(lineas) - 1:
                    p.add_run().add_break()
            sombrear(p, "F3F2EF", es_parrafo=True)
            p.paragraph_format.space_after = Pt(10)
        elif tipo == "tabla":
            tabla = d.add_table(rows=len(dato), cols=len(dato[0]))
            tabla.style = "Table Grid"
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, fila in enumerate(dato):
                for j, valor in enumerate(fila):
                    celda = tabla.cell(i, j)
                    par = celda.paragraphs[0]
                    par.paragraph_format.first_line_indent = Cm(0)
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(0)
                    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = par.add_run(str(valor))
                    run.font.size = Pt(9)
                    run.bold = i == 0
                    if i == 0:
                        sombrear(celda, "E6EEFA")
        elif tipo == "nota":
            d.add_paragraph(dato, style="Nota")
        elif tipo == "figura":
            archivo, leyenda, ancho_cm = [x.strip() for x in dato.split("|")]
            p = d.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(SALIDA / archivo), width=Cm(float(ancho_cm)))
            d.add_paragraph(leyenda, style="Nota")
    d.core_properties.author = AUTOR
    d.core_properties.title = TITULO
    d.save(DOCX)


def main():
    consultas = leer_consultas()
    conexion = mysql.connector.connect(user="root", host="127.0.0.1", database="vinos_tp3",
                                       charset="utf8mb4", collation="utf8mb4_0900_as_ci")
    cur = conexion.cursor()
    bloques, resultados = resolver(parsear(CONTENIDO.read_text(encoding="utf-8")), consultas, cur)
    conexion.close()
    graficos(resultados)
    generar_pdf(bloques)
    generar_docx(bloques)
    print("listo:", PDF.name, DOCX.name)


if __name__ == "__main__":
    main()
